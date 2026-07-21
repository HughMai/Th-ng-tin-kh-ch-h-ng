"""Smoke test: a job with phụ kiện but no cửa.

Regression guard for the A.Của - Công ty Makita case, where a khách was created
but the đơn hàng existed nowhere. Selling only phụ kiện (khóa, bình tích điện,
chi phí khác — no cửa) is ordinary work, and three separate things broke on it:

  1. the tiếp-nhận wizard only created a báo giá when a cửa priced, so the phụ
     kiện, đặt cọc, ngày lắp and ghi chú were all silently discarded
  2. create_quote_header() never recomputed, so such a quote read 0đ — which
     also tripped the "chưa có hạng mục" guard and made chốt impossible
  3. snapshot_order_items_from_quote() added its legacy lump line on top of the
     phụ kiện lines it had just inserted, billing the accessories twice

Also pins the two neighbours that must NOT change: a quote with cửa + phụ kiện,
and a legacy quick quote (lump value_vnd, nothing itemised).

Runs against a throwaway temp SQLite DB — never touches data/htp.db. Env vars
must be set BEFORE `import app`.

Run with: .venv/Scripts/python.exe tests_smoke_phukien_only.py
"""
import os
import tempfile
from pathlib import Path

_tmp_dir = tempfile.mkdtemp(prefix="htp_crm_test_pk_")
os.environ["DB_PATH"] = str(Path(_tmp_dir) / "test.db")
os.environ["FAMILY_PASSWORD"] = "test-pass-123"
os.environ["SESSION_SECRET"] = "test-session-secret"
os.environ["COOKIE_INSECURE"] = "1"

from fastapi.testclient import TestClient

import app
import store

VAT = 0.1
client = TestClient(app.app)
assert client.post("/login", data={"password": "test-pass-123"},
                   follow_redirects=False).status_code == 303

ACC = "Khóa ngang (tole) x2, Bình Tích Điện x1"


def intake(name, phone, units="[]", acc="", deposit="", install="", note=""):
    r = client.post("/khach/tiep-nhan", data={
        "name": name, "phone": phone, "type": "KH", "source": "khac", "units": units,
        "accessories": acc, "deposit": deposit, "install_date": install, "quote_note": note,
    }, follow_redirects=False)
    assert r.status_code == 303, f"intake -> {r.status_code}"
    return r.headers["location"]


# ---- 1. phụ-kiện-only intake creates a real báo giá, keeping every field -----
loc = intake("A.Của - Công ty Makita", "0888981884", acc=ACC,
             deposit="2.000.000", install="2026-07-25", note="Chỉ lắp phụ kiện")
assert loc.startswith("/bao-gia/"), f"phụ-kiện-only intake made no báo giá (went to {loc})"
qid = int(loc.rsplit("/", 1)[1])

q = store.get_quote(qid)
assert q["accessories"] == ACC, f"phụ kiện lost: {q['accessories']!r}"
assert q["deposit_vnd"] == 2_000_000, f"đặt cọc lost: {q['deposit_vnd']}"
assert q["install_date"] == "2026-07-25", f"ngày lắp lost: {q['install_date']}"
assert q["note"] == "Chỉ lắp phụ kiện", f"ghi chú lost: {q['note']}"

# ---- 2. it carries its real total, not the 0đ the INSERT starts with --------
expected = store.pricing.calc_phukien(ACC, [], "KH")
assert expected > 0, "test fixture prices to 0 — pick real phụ kiện"
assert q["value_vnd"] == expected, f"quote value {q['value_vnd']} != phụ kiện total {expected}"

# ---- 3. it is reachable from the pages the family actually looks at ---------
assert "A.Của" in client.get("/bao-gia").text, "missing from /bao-gia"
assert "A.Của" in client.get("/khach").text, "missing from /khach"
assert client.get(f"/bao-gia/{qid}").status_code == 200, "quote page won't render"

# ---- 3b. the save button exists while the quote is still open --------------
# Checked before chốt, which locks the quote and legitimately hides it.
assert "Xong, lưu báo giá" in client.get(f"/bao-gia/{qid}").text, \
    "no 'Xong, lưu báo giá' button on an open phụ-kiện-only báo giá"

# ---- 4. chốt works and bills the phụ kiện exactly once ----------------------
r = client.post(f"/bao-gia/{qid}/trang-thai", data={"trang_thai": "won"},
                follow_redirects=False)
assert r.status_code == 303, f"chốt blocked on a phụ-kiện-only quote: {r.status_code}"
oid = store.get_quote(qid)["order_id"]
assert oid, "chốt created no đơn hàng"

order = store.get_order(oid)
want = expected + round(expected * VAT)
assert order["value_vnd"] == want, \
    f"đơn hàng {order['value_vnd']} != {want} — phụ kiện billed twice?"
lines = store.order_items_for(oid)
assert len(lines) == 2, f"expected 2 phụ kiện lines, got {[l.get('description') for l in lines]}"
assert sum(l["thanh_tien"] for l in lines) == expected, "invoice lines don't sum to the total"

assert "A.Của" in client.get("/don-hang").text, "đơn hàng missing from /don-hang"
assert client.get(f"/don-hang/{oid}").status_code == 200, "order page won't render"

# ---- 4b. the page offers the same actions as any other báo giá -------------
# Without these the family can record a phụ-kiện job but never save it as an
# đơn hàng or send the customer a quote — the buttons were gated on cửa.
page = client.get(f"/bao-gia/{qid}").text
assert "Xuất Báo Giá (Excel)" in page, "no Excel export on a phụ-kiện-only báo giá"
assert "Cửa (0 hạng mục)" not in page, "summary shows a bogus empty cửa line"
assert "Bình Tích Điện" in page, "phụ kiện breakdown missing from the summary"
board = client.get("/bao-gia").text
assert f"/bao-gia/{qid}/xuat" in board, "no Xuất link on the board card"

# ---- 4c. the export renders, and skips the empty cửa table -----------------
r = client.get(f"/bao-gia/{qid}/xuat")
assert r.status_code == 200, f"export blocked on a phụ-kiện-only quote: {r.status_code}"
assert len(r.content) > 4000, "export produced a suspiciously small file"
import io
from openpyxl import load_workbook
sections = [str(x) for row in load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True)
            for x in row if x and ("HẠNG MỤC CỬA" in str(x) or "PHỤ KIỆN" in str(x))]
assert sections == ["I. PHỤ KIỆN"], \
    f"phụ-kiện-only export should be one section numbered I., got {sections}"

# ---- 5. cửa + phụ kiện is unchanged (no double count there either) ----------
units = ('[{"product":"cua_cuon","cong_nghe":"Cửa cuốn công nghệ Úc",'
         '"mau":"Tole màu 5.2 zem","mau_sac":"xanh","ngang":3000,"cao":2200}]')
loc2 = intake("Khách Cửa + PK", "0900111222", units=units, acc="Bình Tích Điện x1")
qid2 = int(loc2.rsplit("/", 1)[1])
q2 = store.get_quote(qid2)
client.post(f"/bao-gia/{qid2}/trang-thai", data={"trang_thai": "won"}, follow_redirects=False)
o2 = store.get_order(store.get_quote(qid2)["order_id"])
assert o2["value_vnd"] == q2["value_vnd"] + round(q2["value_vnd"] * VAT), \
    f"cửa+phụ kiện order {o2['value_vnd']} != quote {q2['value_vnd']} + VAT"

# ---- 6. legacy quick quote (lump value, nothing itemised) still lumps -------
cid = store.create_customer("Khách Gộp", "0900333444", "KH")
qid3 = store.create_quote(cid, "cua_cuon", "Đơn gộp", 5_000_000)
store.set_quote_status(qid3, "won")
oid3 = store.create_order_from_quote(qid3)
o3 = store.get_order(oid3)
assert o3["value_vnd"] == 5_500_000, f"legacy lump order broke: {o3['value_vnd']}"
assert len(store.order_items_for(oid3)) == 1, "legacy lump should stay one generic line"

# ---- 7. a truly empty intake still makes no junk 0đ quote ------------------
loc4 = intake("Khách Trống", "0900555666")
assert loc4.startswith("/khach/"), f"empty intake created a junk báo giá: {loc4}"

print("OK — tests_smoke_phukien_only: phụ-kiện-only intake keeps every field, prices, "
      "chốts into an đơn hàng billed once; cửa+phụ kiện, legacy lump and empty intake unchanged")
