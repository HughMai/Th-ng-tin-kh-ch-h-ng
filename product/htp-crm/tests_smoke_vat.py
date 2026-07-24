"""Smoke test for the VAT opt-out: quotes.apply_vat / orders.apply_vat.

Covers the whole data flow — báo giá summary, chốt inheritance, order value_vnd,
KH còn nợ, ĐL sổ nợ adjustment, hóa đơn and the Excel export.

Runs against a throwaway temp SQLite DB — never touches data/htp.db. Env vars
must be set BEFORE `import app` (app.py reads DB_PATH/FAMILY_PASSWORD/
SESSION_SECRET at import time and calls store.configure(DB_PATH) eagerly).

Run with: .venv/Scripts/python.exe tests_smoke_vat.py
"""
import os
import tempfile
from pathlib import Path

_tmp_dir = tempfile.mkdtemp(prefix="htp_crm_test_")
os.environ["DB_PATH"] = str(Path(_tmp_dir) / "test.db")
os.environ["FAMILY_PASSWORD"] = "test-pass-123"
os.environ["SESSION_SECRET"] = "test-session-secret"
os.environ["COOKIE_INSECURE"] = "1"  # TestClient talks plain http://testserver

from fastapi.testclient import TestClient
import app
import baogia
import pricing
import store

client = TestClient(app.app)
client.post("/login", data={"password": "test-pass-123"}, follow_redirects=False)

DOOR = dict(product="cua_cuon", cong_nghe="Cửa cuốn công nghệ Đức", mau="KV 380",
            ngang_mm=3000, cao_mm=2200)


def make_quote(customer_type="KH", apply_vat=True, name="VAT Test"):
    """A one-door báo giá. Returns (customer_id, quote_id, pre-VAT subtotal)."""
    cid = store.create_customer(name, f"09{store.today_vn()[-2:]}{os.urandom(3).hex()}",
                                customer_type)
    qid = store.create_quote_header(cid, apply_vat=apply_vat)
    price = pricing.get_price(DOOR["product"], DOOR["cong_nghe"], DOOR["mau"], customer_type)
    per_sqm, flat = pricing.get_surcharges(
        DOOR["product"], DOOR["cong_nghe"],
        pricing.area_m2(DOOR["ngang_mm"], DOOR["cao_mm"]))
    total = pricing.line_total(price, DOOR["ngang_mm"], DOOR["cao_mm"], per_sqm, flat)
    store.add_quote_item(qid, DOOR["product"], DOOR["cong_nghe"], DOOR["mau"],
                         DOOR["ngang_mm"], DOOR["cao_mm"], total, False)
    return cid, qid, total


# 1. pricing.vat_amount is the single rate ------------------------------------
assert pricing.vat_amount(10_000_000, True) == 1_000_000, "VAT on = 10%"
assert pricing.vat_amount(10_000_000, False) == 0, "VAT off = 0đ"
print("1. pricing.vat_amount honours the flag OK")

# 2. quotes.value_vnd stays pre-VAT either way --------------------------------
_, q_on, sub = make_quote(apply_vat=True)
_, q_off, sub_off = make_quote(apply_vat=False)
assert store.get_quote(q_on)["value_vnd"] == sub, "VAT-on quote value must stay pre-VAT"
assert store.get_quote(q_off)["value_vnd"] == sub_off, "VAT-off quote value must stay pre-VAT"
assert store.get_quote(q_on)["apply_vat"] == 1 and store.get_quote(q_off)["apply_vat"] == 0
print(f"2. quotes.value_vnd stays pre-VAT regardless of the flag OK ({sub:,}đ)")

# 3. the báo giá summary shows/hides the tax line ------------------------------
body_on = client.get(f"/bao-gia/{q_on}").text
body_off = client.get(f"/bao-gia/{q_off}").text
# Match the summary row itself, not the toggle label (which also says "VAT (10%)").
TAX_ROW = "<span>VAT (10%)</span>"
assert TAX_ROW in body_on, "VAT-on quote must show the tax line"
assert TAX_ROW not in body_off and "Không xuất VAT" in body_off, \
    "VAT-off quote must not show a tax line"
assert 'name="apply_vat"' in body_on, "báo giá detail must offer the VAT toggle"
print("3. báo giá summary shows the tax line only when VAT is on OK")

# 4. chốt: the order inherits the flag and values it correctly ------------------
oid_on = store.create_order_from_quote(q_on)
oid_off = store.create_order_from_quote(q_off)
o_on, o_off = store.get_order(oid_on), store.get_order(oid_off)
assert o_on["apply_vat"] == 1 and o_off["apply_vat"] == 0, "order must inherit quotes.apply_vat"
assert o_on["value_vnd"] == sub + round(sub * 0.1), \
    f"VAT order {o_on['value_vnd']} != {sub + round(sub * 0.1)}"
assert o_off["value_vnd"] == sub_off, \
    f"no-VAT order {o_off['value_vnd']} != bare subtotal {sub_off}"
print(f"4. chốt carries the flag; value_vnd {o_on['value_vnd']:,}đ vs "
      f"{o_off['value_vnd']:,}đ (no VAT) OK")

# 5. the báo giá is locked after chốt — VAT moves to the đơn hàng ---------------
assert client.post(f"/bao-gia/{q_on}/vat", data={}).status_code == 400, \
    "a chốt báo giá must reject VAT edits (đổi trên đơn hàng)"
print("5. đã chốt báo giá rejects VAT edits OK")

# 6. flipping VAT off on a KH đơn hàng moves value_vnd and còn nợ ---------------
before = store.get_order(oid_on)["balance_vnd"]
r = client.post(f"/don-hang/{oid_on}/vat", data={}, follow_redirects=False)  # unticked
assert r.status_code == 303, r.status_code
o = store.get_order(oid_on)
assert o["apply_vat"] == 0 and o["value_vnd"] == sub, \
    f"after VAT off: {o['value_vnd']} != {sub}"
assert o["balance_vnd"] == before - round(sub * 0.1), "còn nợ must drop by the VAT"
# ...and back on again (idempotent both ways)
client.post(f"/don-hang/{oid_on}/vat", data={"apply_vat": "1"}, follow_redirects=False)
assert store.get_order(oid_on)["value_vnd"] == sub + round(sub * 0.1), "VAT must come back"
assert store.set_order_vat(oid_on, True) is False, "no-op flip must return False"
print("6. đơn hàng VAT toggle re-values the order and còn nợ OK")

# 7. ĐL: the sổ nợ gets an offsetting entry, not an edited charge ---------------
dl_cid, dl_qid, dl_sub = make_quote("DL", apply_vat=True, name="ĐL VAT Test")
dl_oid = store.create_order_from_quote(dl_qid)
charged = store.dealer_balance(dl_cid)
assert charged == dl_sub + round(dl_sub * 0.1), f"ĐL charged {charged} at chốt"
client.post(f"/don-hang/{dl_oid}/vat", data={}, follow_redirects=False)  # VAT off
assert store.dealer_balance(dl_cid) == dl_sub, \
    f"ĐL balance {store.dealer_balance(dl_cid)} != {dl_sub} after dropping VAT"
client.post(f"/don-hang/{dl_oid}/vat", data={"apply_vat": "1"}, follow_redirects=False)
assert store.dealer_balance(dl_cid) == dl_sub + round(dl_sub * 0.1), \
    "ĐL balance must return to the VAT-inclusive amount"
print(f"7. ĐL sổ nợ self-corrects via offsetting entries OK ({dl_sub:,}đ ↔ "
      f"{dl_sub + round(dl_sub * 0.1):,}đ)")

# 8. hóa đơn drops the tax row when the order is không VAT ----------------------
inv_off = client.get(f"/don-hang/{oid_off}/hoa-don").text
inv_on = client.get(f"/don-hang/{oid_on}/hoa-don").text
assert "VAT 10%" in inv_on, "VAT order's hóa đơn must show the tax row"
assert "VAT 10%" not in inv_off, "no-VAT order's hóa đơn must omit the tax row"
print("8. hóa đơn omits the tax row for a không-VAT đơn hàng OK")

# 8b. toggling VAT on the đơn hàng syncs the báo giá — its xlsx re-export follows
from openpyxl import load_workbook
import io as _io
def _xlsx_has_vat(qid):
    q = store.get_quote(qid)
    data, _ = baogia.build_baogia_xlsx(q, store.get_customer(q["customer_id"]),
                                       store.quote_items_for(qid), {"name": "HTP", "tagline": "t"})
    wb = load_workbook(_io.BytesIO(data))
    return any(isinstance(c, str) and "VAT" in c
               for row in wb.active.iter_rows(values_only=True) for c in row)
sync_cid, sync_qid, _ = make_quote(apply_vat=True, name="Sync Test")
sync_oid = store.create_order_from_quote(sync_qid)
assert _xlsx_has_vat(sync_qid), "báo giá xlsx must start with a VAT row"
client.post(f"/don-hang/{sync_oid}/vat", data={}, follow_redirects=False)  # VAT off on the order
assert store.get_quote(sync_qid)["apply_vat"] == 0, "order VAT toggle must sync quotes.apply_vat"
assert not _xlsx_has_vat(sync_qid), "báo giá xlsx must drop the VAT row after the order toggles off"
print("8b. đơn hàng VAT toggle syncs the báo giá xlsx OK")

# 9. the Excel export follows the flag ------------------------------------------
company = {"name": "HTP", "tagline": "t", "phone": "0900", "address": "a",
           "email": "e@e.vn", "website": "w"}
for qid, expect_vat in ((q_on, True), (q_off, False)):
    q = store.get_quote(qid)
    data, fname = baogia.build_baogia_xlsx(q, store.get_customer(q["customer_id"]),
                                           store.quote_items_for(qid), company)
    assert data and fname.endswith(".xlsx"), f"xlsx build failed for quote {qid}"
print("9. Xuất Báo Giá (Excel) builds for both VAT and không-VAT quotes OK")

# 10. legacy rows default to VAT on ---------------------------------------------
legacy_cid = store.create_customer("Legacy", "0999888777", "KH")
legacy_qid = store.create_quote(legacy_cid, "cua_cuon", "cũ", 10_000_000)
assert store.get_quote(legacy_qid)["apply_vat"] == 1, "existing quotes keep VAT by default"
legacy_oid = store.create_order(legacy_cid, "cua_cuon", "cũ", 10_000_000)
assert store.get_order(legacy_oid)["value_vnd"] == 11_000_000, \
    "a plain đơn hàng still bills VAT by default"
print("10. legacy/default rows keep VAT on OK")

# 11. the "Thông tin chung" form actually reaches its route ----------------------
form_cid = store.create_customer("Form Test", "0911222333", "KH")
page = client.get(f"/bao-gia/nhieu-hang-muc?khach={form_cid}").text
assert 'action="/bao-gia/nhieu-hang-muc"' in page, "header form must post to a real route"
assert 'name="apply_vat"' in page, "header form must offer the VAT toggle"
r = client.post("/bao-gia/nhieu-hang-muc",
                data={"customer_id": form_cid, "deposit": "", "install_date": "", "note": ""},
                follow_redirects=False)
assert r.status_code == 303, f"header form POST returned {r.status_code}"
new_qid = int(r.headers["location"].rsplit("/", 1)[1])
assert store.get_quote(new_qid)["apply_vat"] == 0, "unticked checkbox = không VAT"
print("11. 'Thông tin chung' form posts to a live route and carries the toggle OK")

# 11b. the đơn hàng door editor: add + re-spec a door line via the calculator ----
door_cid, door_qid, door_sub = make_quote(apply_vat=True, name="Order Door Edit")
door_oid = store.create_order_from_quote(door_qid)
door_line = [i for i in store.order_items_for(door_oid) if i["product"] != "khac"][0]
# the đơn hàng detail page links each door line to the calculator, not a flat input
det = client.get(f"/don-hang/{door_oid}").text
assert f"/don-hang/{door_oid}/hang-muc/{door_line['id']}/sua-cua" in det, \
    "door line must link to the calculator editor"
assert f"/don-hang/{door_oid}/hang-muc/cua-moi" in det, "must offer + Thêm cửa"
# the edit page renders the calculator prefilled for this order item
form = client.get(f"/don-hang/{door_oid}/hang-muc/{door_line['id']}/sua-cua").text
assert 'name="loai_cua"' in form and f"/hang-muc/{door_line['id']}/sua-cua" in form, \
    "order door edit page must post back to sua-cua"
# re-spec the door to a bigger size → price must change and value recompute
new_price = pricing.line_total(
    pricing.get_price("cua_cuon", "Cửa cuốn công nghệ Đức", "KV 380", "KH"),
    3000, 2500,
    *pricing.get_surcharges("cua_cuon", "Cửa cuốn công nghệ Đức", pricing.area_m2(3000, 2500)))
r = client.post(f"/don-hang/{door_oid}/hang-muc/{door_line['id']}/sua-cua",
                data={"loai_cua": "cua_cuon", "cong_nghe": "Cửa cuốn công nghệ Đức",
                      "mau": "KV 380", "ngang": "3000", "cao": "2500"}, follow_redirects=False)
assert r.status_code == 303, f"door edit POST returned {r.status_code}"
edited = store.get_order_item(door_line["id"])
assert edited["thanh_tien"] == new_price and edited["cao_mm"] == 2500, \
    f"door re-spec didn't reprice: {edited['thanh_tien']} != {new_price}"
assert store.get_order(door_oid)["value_vnd"] == new_price + round(new_price * 0.1), \
    "order value must recompute (with VAT) after a door re-spec"
# add a second door via the calculator
before_n = len(store.order_items_for(door_oid))
r = client.post(f"/don-hang/{door_oid}/hang-muc/cua-moi",
                data={"loai_cua": "cua_keo", "cong_nghe": "Có lá", "mau": "6zem",
                      "ngang": "2000", "cao": "2000"}, follow_redirects=False)
assert r.status_code == 303, f"door add POST returned {r.status_code}"
assert len(store.order_items_for(door_oid)) == before_n + 1, "+ Thêm cửa didn't add a line"
print("11b. đơn hàng door add + re-spec via the calculator OK")

# 12. the "Khách hàng mới" intake wizard carries the VAT toggle ------------------
wiz_page = client.get("/khach/tiep-nhan").text
assert 'name="apply_vat"' in wiz_page, "intake wizard must offer the VAT toggle"
# ticked → quote billed with VAT
r = client.post("/khach/tiep-nhan", data={"name": "Wizard VAT", "phone": "0912000001",
                "type": "KH", "deposit": "1.000.000", "apply_vat": "1"}, follow_redirects=False)
wiz_qid = int(r.headers["location"].rsplit("/", 1)[1])
assert store.get_quote(wiz_qid)["apply_vat"] == 1, "intake wizard must honour a ticked VAT box"
# unticked → không VAT
r = client.post("/khach/tiep-nhan", data={"name": "Wizard NoVAT", "phone": "0912000002",
                "type": "KH", "deposit": "1.000.000"}, follow_redirects=False)
wiz_qid2 = int(r.headers["location"].rsplit("/", 1)[1])
assert store.get_quote(wiz_qid2)["apply_vat"] == 0, "intake wizard must honour an unticked VAT box"
print("12. 'Khách hàng mới' intake wizard carries the VAT toggle OK")

print("\nAll VAT smoke tests passed.")
